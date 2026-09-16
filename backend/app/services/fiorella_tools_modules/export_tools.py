from __future__ import annotations

import os
import re

from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font

from app.services.fiorella_tools import (
    tool_search_punches,
)


# ============================================================
# DIRECTORIO DE EXPORTACIONES
# ============================================================

EXPORT_DIR = Path(
    os.getenv(
        "FIORELLA_EXPORT_DIR",
        "data/fiorella_exports",
    )
)

EXPORT_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


# ============================================================
# NOMBRE SEGURO
# ============================================================

def _safe_filename(
    value: str,
) -> str:

    value = re.sub(
        r"[^a-zA-Z0-9_-]+",
        "_",
        value,
    )

    return value.strip("_")


# ============================================================
# EXPORTAR PONCHES A EXCEL
# ============================================================

def export_punches_excel(
    user: dict[str, Any],
    *,
    query: str | None = None,
    fecha_desde: str | None = None,
    fecha_hasta: str | None = None,
    dispositivo: str | None = None,
    limit: int = 5000,
) -> dict[str, Any]:

        # ========================================================
    # NORMALIZAR ARGUMENTOS RECIBIDOS DESDE EL MODELO
    # ========================================================

    normalized_query = (
        str(query).strip()
        if query is not None
        else None
    )

    normalized_device = (
        str(dispositivo).strip()
        if dispositivo is not None
        else None
    )

    normalized_limit = min(
        max(
            int(limit or 5000),
            1,
        ),
        5000,
    )

    # --------------------------------------------------------
    # Obtener datos usando la tool existente
    # --------------------------------------------------------

    result = tool_search_punches(
        user,
        query=normalized_query,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        dispositivo=normalized_device,
        limit=normalized_limit,
    )

    if not result.get(
        "ok",
        False,
    ):
        return result


    # --------------------------------------------------------
    # Detectar colección de registros
    # --------------------------------------------------------

    rows = (
        result.get("items")
        or result.get("rows")
        or result.get("data")
        or result.get("resultados")
        or []
    )


    if not rows:

        return {
            "ok": False,
            "error": (
                "No existen registros "
                "para exportar."
            ),
        }


    if not isinstance(
        rows,
        list,
    ):

        return {
            "ok": False,
            "error": (
                "La consulta devolvió "
                "un formato no compatible."
            ),
        }


    first_row = rows[0]


    if not isinstance(
        first_row,
        dict,
    ):

        return {
            "ok": False,
            "error": (
                "Formato de registros "
                "no compatible con Excel."
            ),
        }


    # --------------------------------------------------------
    # Nombre del archivo
    # --------------------------------------------------------

    timestamp = datetime.now().strftime(
        "%Y%m%d_%H%M%S"
    )


    parts = [
        "ponches",
        query or "todos",
        fecha_desde or "",
        fecha_hasta or "",
        timestamp,
    ]


    filename = (
        _safe_filename(
            "_".join(
                part
                for part in parts
                if part
            )
        )
        + ".xlsx"
    )


    filepath = (
        EXPORT_DIR
        / filename
    )


    # --------------------------------------------------------
    # Crear Excel
    # --------------------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        "Ponches"
    )


    headers = list(
        first_row.keys()
    )


    # --------------------------------------------------------
    # Encabezados
    # --------------------------------------------------------

    for column_index, header in enumerate(
        headers,
        start=1,
    ):

        cell = worksheet.cell(
            row=1,
            column=column_index,
            value=header,
        )

        cell.font = Font(
            bold=True,
        )


    # --------------------------------------------------------
    # Registros
    # --------------------------------------------------------

    for row_index, row in enumerate(
        rows,
        start=2,
    ):

        if not isinstance(
            row,
            dict,
        ):
            continue


        for column_index, header in enumerate(
            headers,
            start=1,
        ):

            value = row.get(
                header
            )


            worksheet.cell(
                row=row_index,
                column=column_index,
                value=(
                    str(value)
                    if value is not None
                    else ""
                ),
            )


    # --------------------------------------------------------
    # Auto ancho de columnas
    # --------------------------------------------------------

    for column_cells in worksheet.columns:

        max_length = 0

        column_letter = (
            column_cells[0]
            .column_letter
        )


        for cell in column_cells:

            value = (
                cell.value
                if cell.value is not None
                else ""
            )

            max_length = max(
                max_length,
                len(
                    str(value)
                ),
            )


        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            40,
        )


    # --------------------------------------------------------
    # Guardar
    # --------------------------------------------------------

    workbook.save(
        filepath
    )


    download_url = (
        f"/api/v1/fiorella/files/"
        f"{filename}"
    )


    # --------------------------------------------------------
    # Resultado
    # --------------------------------------------------------

    return {
        "ok": True,

        "count": len(
            rows
        ),

        "filename": filename,

        "download_url": (
            download_url
        ),

        "action": {
            "type": "download",
            "url": download_url,
            "filename": filename,
        },

        "message": (
            f"Generé el archivo Excel "
            f"con {len(rows)} registros."
        ),
    }