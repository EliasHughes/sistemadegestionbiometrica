from datetime import date, datetime, time
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query, Request

from app.core.deps import get_current_user
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from app.services.database import fetch_all, test_connection


router = APIRouter(
    prefix="/records/export",
    tags=["export"],
    dependencies=[Depends(get_current_user)],
)


# ============================================================
# UTILIDADES
# ============================================================

def _cell_value(value):
    """Convierte valores de fecha/hora para Excel sin romper datos nulos."""
    if value is None:
        return ""

    if isinstance(value, (datetime, date, time)):
        return str(value)

    return value


def _clean_string(value: Optional[str]) -> Optional[str]:
    """Limpia strings recibidos desde query parameters."""
    if value is None:
        return None

    value = str(value).strip()

    return value if value else None


def _parse_date(value: Optional[str], field_name: str) -> Optional[date]:
    """
    Convierte YYYY-MM-DD a date.

    El input HTML type="date" envía exactamente este formato.
    """
    value = _clean_string(value)

    if not value:
        return None

    try:
        return date.fromisoformat(value)
    except ValueError:
        raise HTTPException(
            status_code=422,
            detail=(
                f"El parámetro '{field_name}' no es válido. "
                f"Debe utilizar el formato YYYY-MM-DD."
            ),
        )


def _resolve_date_range(
    fecha: Optional[str] = None,
    fecha_desde: Optional[str] = None,
    fecha_hasta: Optional[str] = None,
    desde: Optional[str] = None,
    hasta: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    fecha_inicio: Optional[str] = None,
    fecha_fin: Optional[str] = None,
) -> tuple[Optional[date], Optional[date]]:
    """
    Resuelve el rango de fechas aceptando tanto el parámetro antiguo
    'fecha' como los nombres utilizados por el frontend.

    Prioridad:
      1. fecha_desde / fecha_hasta
      2. desde / hasta
      3. start_date / end_date
      4. fecha_inicio / fecha_fin
      5. fecha como filtro de un solo día

    Esto permite compatibilidad con versiones anteriores del frontend.
    """

    raw_desde = (
        _clean_string(fecha_desde)
        or _clean_string(desde)
        or _clean_string(start_date)
        or _clean_string(fecha_inicio)
    )

    raw_hasta = (
        _clean_string(fecha_hasta)
        or _clean_string(hasta)
        or _clean_string(end_date)
        or _clean_string(fecha_fin)
    )

    # Compatibilidad con el endpoint original:
    # ?fecha=2026-09-05 significa exactamente ese día.
    if not raw_desde and not raw_hasta:
        single = _clean_string(fecha)

        if single:
            parsed = _parse_date(single, "fecha")
            return parsed, parsed

        return None, None

    parsed_desde = _parse_date(raw_desde, "fecha_desde")
    parsed_hasta = _parse_date(raw_hasta, "fecha_hasta")

    if parsed_desde and parsed_hasta and parsed_desde > parsed_hasta:
        raise HTTPException(
            status_code=422,
            detail="La fecha 'Desde' no puede ser posterior a la fecha 'Hasta'.",
        )

    return parsed_desde, parsed_hasta


def _resolve_device(
    dispositivo: Optional[str] = None,
    device: Optional[str] = None,
    reloj: Optional[str] = None,
) -> Optional[str]:
    """
    Acepta los tres nombres utilizados por distintas versiones
    del frontend.
    """
    value = _clean_string(dispositivo) or _clean_string(device) or _clean_string(reloj)

    if not value:
        return None

    if value.lower() == "todos":
        return None

    return value


# ============================================================
# CONSULTA CENTRALIZADA DE PONCHES
# ============================================================

def _query_punches(
    limit: int,
    fecha: Optional[str] = None,
    dispositivo: Optional[str] = None,
    *,
    fecha_desde: Optional[str] = None,
    fecha_hasta: Optional[str] = None,
    desde: Optional[str] = None,
    hasta: Optional[str] = None,
    start_date: Optional[str] = None,
    end_date: Optional[str] = None,
    fecha_inicio: Optional[str] = None,
    fecha_fin: Optional[str] = None,
    device: Optional[str] = None,
    reloj: Optional[str] = None,
):
    """
    Consulta los ponches aplicando realmente el rango seleccionado.

    IMPORTANTE:
    No se utiliza:
        fecha = ?

    porque eso falla cuando la columna 'fecha' es DATETIME y/o
    cuando el frontend envía un rango.

    En su lugar se utiliza:

        fecha >= fecha_desde
        AND fecha < fecha_hasta + 1 día

    Esto incluye TODO el día seleccionado, hasta 23:59:59.999999,
    sin depender de la precisión de SQL Server.
    """

    db = test_connection()

    if db["status"] != "online":
        raise HTTPException(
            status_code=503,
            detail=db.get("detail", "La base de datos no está disponible."),
        )

    # --------------------------------------------------------
    # Resolver fechas
    # --------------------------------------------------------
    date_from, date_to = _resolve_date_range(
        fecha=fecha,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        desde=desde,
        hasta=hasta,
        start_date=start_date,
        end_date=end_date,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
    )

    # --------------------------------------------------------
    # Resolver dispositivo
    # --------------------------------------------------------
    device_value = _resolve_device(
        dispositivo=dispositivo,
        device=device,
        reloj=reloj,
    )

    conditions = []
    params = []

    # --------------------------------------------------------
    # FILTRO DE FECHA
    # --------------------------------------------------------
    #
    # Si solo se selecciona Desde:
    #
    #     fecha >= Desde
    #
    # Si solo se selecciona Hasta:
    #
    #     fecha < Hasta + 1 día
    #
    # Si se seleccionan ambas:
    #
    #     fecha >= Desde
    #     fecha < Hasta + 1 día
    #
    # Este último punto es especialmente importante:
    #
    # Desde 05/09/2026
    # Hasta 05/09/2026
    #
    # devuelve exclusivamente los registros del 05/09/2026.
    # --------------------------------------------------------

    if date_from:
        conditions.append("fecha >= ?")
        params.append(date_from)

    if date_to:
        conditions.append("fecha < DATEADD(day, 1, ?)")
        params.append(date_to)

    # --------------------------------------------------------
    # FILTRO DE DISPOSITIVO
    # --------------------------------------------------------
    if device_value:
        conditions.append(
            "LTRIM(RTRIM(dispositivo_origen)) = LTRIM(RTRIM(?))"
        )
        params.append(device_value)

    where_clause = ""

    if conditions:
        where_clause = "WHERE " + " AND ".join(conditions)

    # TOP (?) también es parametrizado.
    sql_params = [int(limit), *params]

    sql = f"""
        SELECT TOP (?)
            codigo,
            nombre,
            departamento,
            fecha,
            entrada,
            salida,
            dispositivo_origen,
            ultima_sincronizacion
        FROM [dbo].[punches]
        {where_clause}
        ORDER BY fecha DESC, entrada DESC
    """

    try:
        rows = fetch_all(sql, tuple(sql_params))

        return rows

    except Exception as exc:
        raise HTTPException(
            status_code=500,
            detail=f"Error SQL al consultar los ponches: {exc}",
        )


# ============================================================
# PARÁMETROS COMPARTIDOS DE EXPORTACIÓN
# ============================================================

def _export_query_parameters(
    limit: int = Query(1000, ge=1, le=50000),
    fecha: Optional[str] = Query(None),

    fecha_desde: Optional[str] = Query(None),
    fecha_hasta: Optional[str] = Query(None),

    desde: Optional[str] = Query(None),
    hasta: Optional[str] = Query(None),

    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),

    fecha_inicio: Optional[str] = Query(None),
    fecha_fin: Optional[str] = Query(None),

    dispositivo: Optional[str] = Query(None),
    device: Optional[str] = Query(None),
    reloj: Optional[str] = Query(None),
):
    """
    Agrupa los parámetros de las exportaciones.

    Se mantienen múltiples nombres para no romper el frontend
    existente.
    """
    return {
        "limit": limit,
        "fecha": fecha,
        "fecha_desde": fecha_desde,
        "fecha_hasta": fecha_hasta,
        "desde": desde,
        "hasta": hasta,
        "start_date": start_date,
        "end_date": end_date,
        "fecha_inicio": fecha_inicio,
        "fecha_fin": fecha_fin,
        "dispositivo": dispositivo,
        "device": device,
        "reloj": reloj,
    }


# ============================================================
# EXCEL
# ============================================================

@router.api_route("/excel", methods=["GET", "POST"])
def export_excel(
    limit: int = Query(1000, ge=1, le=50000),
    fecha: Optional[str] = Query(None),

    fecha_desde: Optional[str] = Query(None),
    fecha_hasta: Optional[str] = Query(None),

    desde: Optional[str] = Query(None),
    hasta: Optional[str] = Query(None),

    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),

    fecha_inicio: Optional[str] = Query(None),
    fecha_fin: Optional[str] = Query(None),

    dispositivo: Optional[str] = Query(None),
    device: Optional[str] = Query(None),
    reloj: Optional[str] = Query(None),
):
    """
    Exporta los ponches filtrados a Excel.
    """

    rows = _query_punches(
        limit=limit,
        fecha=fecha,
        dispositivo=dispositivo,

        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,

        desde=desde,
        hasta=hasta,

        start_date=start_date,
        end_date=end_date,

        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,

        device=device,
        reloj=reloj,
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Ponches"

    headers = [
        "Código",
        "Nombre",
        "Departamento",
        "Fecha",
        "Entrada",
        "Salida",
        "Dispositivo",
        "Última sincronización",
    ]

    header_fill = PatternFill(
        "solid",
        fgColor="C8102E",
    )

    header_font = Font(
        color="FFFFFF",
        bold=True,
        size=11,
    )

    thin = Border(
        left=Side(style="thin", color="DDDDDD"),
        right=Side(style="thin", color="DDDDDD"),
        top=Side(style="thin", color="DDDDDD"),
        bottom=Side(style="thin", color="DDDDDD"),
    )

    # --------------------------------------------------------
    # ENCABEZADOS
    # --------------------------------------------------------

    for col, header in enumerate(headers, 1):
        cell = ws.cell(1, col, header)

        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = thin

    # --------------------------------------------------------
    # DATOS
    # --------------------------------------------------------

    for row_index, row in enumerate(rows, 2):

        values = [
            row.get("codigo"),
            row.get("nombre"),
            row.get("departamento"),
            row.get("fecha"),
            row.get("entrada"),
            row.get("salida"),
            row.get("dispositivo_origen"),
            row.get("ultima_sincronizacion"),
        ]

        for col_index, value in enumerate(values, 1):

            cell = ws.cell(
                row_index,
                col_index,
                _cell_value(value),
            )

            cell.border = thin

    # --------------------------------------------------------
    # ANCHO DE COLUMNAS
    # --------------------------------------------------------

    widths = [
        12,
        22,
        16,
        14,
        10,
        10,
        28,
        22,
    ]

    letters = [
        "A",
        "B",
        "C",
        "D",
        "E",
        "F",
        "G",
        "H",
    ]

    for column_letter, width in zip(letters, widths):
        ws.column_dimensions[column_letter].width = width

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # --------------------------------------------------------
    # ARCHIVO
    # --------------------------------------------------------

    buffer = BytesIO()

    wb.save(buffer)
    buffer.seek(0)

    device_value = (
        _resolve_device(
            dispositivo=dispositivo,
            device=device,
            reloj=reloj,
        )
        or "todos"
    )

    date_from, date_to = _resolve_date_range(
        fecha=fecha,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        desde=desde,
        hasta=hasta,
        start_date=start_date,
        end_date=end_date,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
    )

    from_tag = date_from.isoformat() if date_from else "inicio"
    to_tag = date_to.isoformat() if date_to else "fin"

    safe_device = (
        device_value
        .replace(" ", "_")
        .replace("/", "_")
        .replace("\\", "_")
        .replace('"', "")
        .replace(":", "_")
        [:40]
    )

    filename = (
        f"ponches_{from_tag}_a_{to_tag}_{safe_device}.xlsx"
    )

    return StreamingResponse(
        buffer,
        media_type=(
            "application/"
            "vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )


# ============================================================
# PDF
# ============================================================

@router.api_route("/pdf", methods=["GET", "POST"])
def export_pdf(
    limit: int = Query(500, ge=1, le=3000),
    fecha: Optional[str] = Query(None),

    fecha_desde: Optional[str] = Query(None),
    fecha_hasta: Optional[str] = Query(None),

    desde: Optional[str] = Query(None),
    hasta: Optional[str] = Query(None),

    start_date: Optional[str] = Query(None),
    end_date: Optional[str] = Query(None),

    fecha_inicio: Optional[str] = Query(None),
    fecha_fin: Optional[str] = Query(None),

    dispositivo: Optional[str] = Query(None),
    device: Optional[str] = Query(None),
    reloj: Optional[str] = Query(None),
):
    """
    Exporta los ponches filtrados a PDF.
    """

    rows = _query_punches(
        limit=limit,
        fecha=fecha,
        dispositivo=dispositivo,

        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,

        desde=desde,
        hasta=hasta,

        start_date=start_date,
        end_date=end_date,

        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,

        device=device,
        reloj=reloj,
    )

    buffer = BytesIO()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "TitleCI",
        parent=styles["Heading1"],
        fontSize=14,
        textColor=colors.HexColor("#C8102E"),
        spaceAfter=6,
    )

    sub_style = ParagraphStyle(
        "SubCI",
        parent=styles["Normal"],
        fontSize=9,
        textColor=colors.HexColor("#52525b"),
        spaceAfter=12,
    )

    # --------------------------------------------------------
    # INFORMACIÓN DEL FILTRO
    # --------------------------------------------------------

    date_from, date_to = _resolve_date_range(
        fecha=fecha,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        desde=desde,
        hasta=hasta,
        start_date=start_date,
        end_date=end_date,
        fecha_inicio=fecha_inicio,
        fecha_fin=fecha_fin,
    )

    device_value = _resolve_device(
        dispositivo=dispositivo,
        device=device,
        reloj=reloj,
    )

    filtro = []

    if date_from and date_to:
        if date_from == date_to:
            filtro.append(
                f"Fecha: {date_from.isoformat()}"
            )
        else:
            filtro.append(
                "Desde: "
                f"{date_from.isoformat()} "
                "Hasta: "
                f"{date_to.isoformat()}"
            )

    elif date_from:
        filtro.append(
            f"Desde: {date_from.isoformat()}"
        )

    elif date_to:
        filtro.append(
            f"Hasta: {date_to.isoformat()}"
        )

    if device_value:
        filtro.append(
            f"Reloj: {device_value}"
        )
    else:
        filtro.append("Reloj: Todos")

    filtro.append(
        f"Registros: {len(rows)}"
    )

    elements = [
        Paragraph(
            "Visualizador de Ponches · César Iglesias",
            title_style,
        ),
        Paragraph(
            " · ".join(filtro),
            sub_style,
        ),
        Spacer(1, 6),
    ]

    # --------------------------------------------------------
    # TABLA
    # --------------------------------------------------------

    data = [[
        "Código",
        "Nombre",
        "Depto",
        "Fecha",
        "Entrada",
        "Salida",
        "Dispositivo",
    ]]

    for row in rows:
        data.append([
            str(row.get("codigo") or ""),
            str(row.get("nombre") or "")[:22],
            str(row.get("departamento") or "")[:12],
            str(row.get("fecha") or "")[:10],
            str(row.get("entrada") or "")[:8],
            str(row.get("salida") or "")[:8],
            str(row.get("dispositivo_origen") or "")[:28],
        ])

    table = Table(
        data,
        repeatRows=1,
        colWidths=[
            70,
            120,
            70,
            70,
            55,
            55,
            160,
        ],
    )

    table.setStyle(
        TableStyle([
            (
                "BACKGROUND",
                (0, 0),
                (-1, 0),
                colors.HexColor("#C8102E"),
            ),
            (
                "TEXTCOLOR",
                (0, 0),
                (-1, 0),
                colors.white,
            ),
            (
                "FONTNAME",
                (0, 0),
                (-1, 0),
                "Helvetica-Bold",
            ),
            (
                "FONTSIZE",
                (0, 0),
                (-1, 0),
                8,
            ),
            (
                "FONTSIZE",
                (0, 1),
                (-1, -1),
                7,
            ),
            (
                "ALIGN",
                (0, 0),
                (-1, 0),
                "CENTER",
            ),
            (
                "GRID",
                (0, 0),
                (-1, -1),
                0.4,
                colors.HexColor("#e4e4e7"),
            ),
            (
                "ROWBACKGROUNDS",
                (0, 1),
                (-1, -1),
                [
                    colors.white,
                    colors.HexColor("#fafafa"),
                ],
            ),
            (
                "VALIGN",
                (0, 0),
                (-1, -1),
                "MIDDLE",
            ),
            (
                "TOPPADDING",
                (0, 0),
                (-1, -1),
                3,
            ),
            (
                "BOTTOMPADDING",
                (0, 0),
                (-1, -1),
                3,
            ),
        ])
    )

    elements.append(table)

    doc.build(elements)

    buffer.seek(0)

    safe_device = (
        (device_value or "todos")
        .replace(" ", "_")
        .replace("/", "_")
        .replace("\\", "_")
        .replace('"', "")
        .replace(":", "_")
        [:40]
    )

    from_tag = (
        date_from.isoformat()
        if date_from
        else "inicio"
    )

    to_tag = (
        date_to.isoformat()
        if date_to
        else "fin"
    )

    filename = (
        f"ponches_{from_tag}_a_{to_tag}_{safe_device}.pdf"
    )

    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={
            "Content-Disposition": (
                f'attachment; filename="{filename}"'
            )
        },
    )
